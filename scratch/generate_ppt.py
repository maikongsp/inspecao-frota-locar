import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # PALETA EXATA E OFICIAL LOCAR GUINDASTES: CINZA, AMARELO E BRANCO
    # Conforme css/main.css do Sistema Locar:
    # - Cinza Fundo: #383535 (RGB 56, 53, 53)
    # - Cinza Superfície Oficial: #474444 (RGB 71, 68, 68)
    # - Cinza Cards e Contraste: #2E2B2B (RGB 46, 43, 43)
    # - Cinza Bordas e Divisores: #5C5656 (RGB 92, 86, 86) / #736C6C (RGB 115, 108, 108)
    # - Amarelo Canário Oficial Locar: #FFF212 (RGB 255, 242, 18)
    # - Amarelo Ouro de Apoio: #F59E0B (RGB 245, 158, 11)
    # - Branco Puro: #FFFFFF (RGB 255, 255, 255)
    # - Branco Acinzentado (Textos): #F1F5F9 (RGB 241, 245, 249) / #CBD5E1 (RGB 203, 213, 225)
    # =========================================================================

    LOCAR_BG_DARK = RGBColor(56, 53, 53)       # #383535 (Cinza Base)
    LOCAR_SURFACE = RGBColor(71, 68, 68)       # #474444 (Cinza Superfície Oficial da Logomarca)
    LOCAR_CARD = RGBColor(46, 43, 43)          # #2E2B2B (Cinza Card Profundo)
    LOCAR_BORDER = RGBColor(92, 86, 86)        # #5C5656 (Borda Cinza)
    LOCAR_BORDER_LIGHT = RGBColor(115, 108, 108) # #736C6C (Borda Suave)
    
    LOCAR_YELLOW = RGBColor(255, 242, 18)      # #FFF212 (Amarelo Oficial Locar)
    LOCAR_YELLOW_WARM = RGBColor(245, 158, 11) # #F59E0B (Amarelo Industrial)
    
    TEXT_WHITE = RGBColor(255, 255, 255)       # Branco Puro
    TEXT_LIGHT = RGBColor(241, 245, 249)       # Branco Leve
    TEXT_MUTED = RGBColor(203, 213, 225)       # Cinza Claro Leitura

    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    assets_logo = os.path.join(base_dir, 'assets', 'locar_logo_transparent.png')
    has_logo = os.path.exists(assets_logo)

    screenshots_dir = os.path.join(base_dir, 'assets', 'screenshots')
    img_dash = os.path.join(screenshots_dir, 'screen_dashboard.png')
    img_insp = os.path.join(screenshots_dir, 'screen_inspecao.png')
    img_pcm = os.path.join(screenshots_dir, 'screen_pcm.png')
    img_laudos = os.path.join(screenshots_dir, 'screen_laudos.png')

    SYSTEM_URL = "https://inspecao-frota-locar.vercel.app/"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LOCAR_BG_DARK
        bg.line.fill.background()
        
        # Faixa industrial amarela oficial Locar
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = LOCAR_YELLOW
        stripe.line.fill.background()

    def add_header(slide, title_text, category_text="LOCAR GUINDASTES E TRANSPORTES INTERMODAIS S/A • FILIAL BETIM/MG"):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.32), Inches(9.5), Inches(1.1))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = LOCAR_YELLOW
        p_cat.font.name = "Arial"
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Arial"
        p_title.space_before = Pt(3)

        if has_logo:
            slide.shapes.add_picture(assets_logo, Inches(10.8), Inches(0.3), width=Inches(1.8))

    def add_card(slide, left, top, width, height, bg_color=LOCAR_CARD, border_color=LOCAR_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.2)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: CAPA INSTITUCIONAL COM LINK DE ACESSO AO SISTEMA
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Detalhe visual amarelo lateral
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.9), Inches(0.12), Inches(3.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = LOCAR_YELLOW
    accent.line.fill.background()

    # Logo oficial
    if has_logo:
        slide1.shapes.add_picture(assets_logo, Inches(0.8), Inches(0.85), width=Inches(2.5))

    # Título Principal e Subtítulo
    title_box = slide1.shapes.add_textbox(Inches(1.15), Inches(1.8), Inches(11.2), Inches(3.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "MANUAL FUNCIONAL E OPERACIONAL • LOCAR GUINDASTES"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW
    p.font.name = "Arial"

    p = tf1.add_paragraph()
    p.text = "Sistema de Inspeção & Qualidade de Frota"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.font.name = "Arial"
    p.space_before = Pt(6)

    p = tf1.add_paragraph()
    p.text = "Gestão Operacional de Pátio, Checklists Normativos (NR-11/12/18/35), Disparo Automático ao PCM,\nSincronização Nuvem Appwrite Cloud, Arquitetura 100% Offline-First e Laudos Periciais com Hash Forense."
    p.font.size = Pt(13.5)
    p.font.color.rgb = TEXT_LIGHT
    p.font.name = "Arial"
    p.space_before = Pt(10)

    # CARD DE DESTAQUE: LINK OFICIAL DE ACESSO AO SISTEMA
    link_card = add_card(slide1, Inches(1.15), Inches(4.3), Inches(11.0), Inches(1.15), LOCAR_SURFACE, LOCAR_YELLOW)
    link_box = slide1.shapes.add_textbox(Inches(1.35), Inches(4.4), Inches(10.6), Inches(0.95))
    tfl = link_box.text_frame
    tfl.word_wrap = True

    p = tfl.paragraphs[0]
    p.text = "🌐 LINK OFICIAL DE ACESSO AO SISTEMA (PRODUÇÃO VERCEL & PWA):"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tfl.add_paragraph()
    p.text = SYSTEM_URL
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(2)

    p = tfl.add_paragraph()
    p.text = "Disponível no Computador, Tablet ou Celular (Instalável via PWA). Funciona 100% Offline no Pátio."
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(2)

    # Rodapé da Capa
    add_card(slide1, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.1), LOCAR_CARD, LOCAR_BORDER)
    tb_foot = slide1.shapes.add_textbox(Inches(1.1), Inches(5.95), Inches(11.1), Inches(0.8))
    tff = tb_foot.text_frame
    tff.word_wrap = True
    p = tff.paragraphs[0]
    p.text = "Unidade Operacional: Locar Guindastes e Transportes Intermodais S/A — Filial Betim / MG"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p = tff.add_paragraph()
    p.text = "Desenvolvido por Maikon Pinho  •  Identidade Visual: Cinza Oficial (#474444), Amarelo (#FFF212) e Branco  •  2026"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: TELAS REAIS DO SISTEMA & NAVEGAÇÃO PRINCIPAL
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "1. Visão Geral da Interface & Telas do Sistema", "EXEMPLOS VISUAIS REAIS DO SISTEMA EM OPERAÇÃO")

    screens_data = [
        {"img": img_dash, "title": "Dashboard da Frota", "sub": "KPIs em tempo real e grid dos 272 ativos de Betim."},
        {"img": img_insp, "title": "Stepper de Inspeção", "sub": "Checklist normativo com fotos obrigatórias."},
        {"img": img_pcm, "title": "Central do PCM", "sub": "S.S. automáticas, ordens e disparos de e-mail/WhatsApp."},
        {"img": img_laudos, "title": "Laudos & Histórico", "sub": "Rastreabilidade e certificação com QR Code / SHA-256."}
    ]

    card_w = Inches(2.78)
    card_h = Inches(4.9)
    for i, s in enumerate(screens_data):
        sx = Inches(0.8 + i * 2.98)
        sy = Inches(1.8)
        add_card(slide2, sx, sy, card_w, card_h, LOCAR_CARD, LOCAR_BORDER)

        # Imagem real do sistema
        if os.path.exists(s["img"]):
            slide2.shapes.add_picture(s["img"], sx + Inches(0.12), sy + Inches(0.15), width=card_w - Inches(0.24))

        # Texto descritivo
        tb = slide2.shapes.add_textbox(sx + Inches(0.15), sy + Inches(3.2), card_w - Inches(0.3), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = s["title"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = LOCAR_YELLOW

        p = tf.add_paragraph()
        p.text = s["sub"]
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 3: OPERAÇÃO OFFLINE-FIRST NO CELULAR
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "2. Funcionamento 100% Offline no Celular (PWA)", "RESILIÊNCIA NO PÁTIO DE BETIM SEM SINAL DE REDE")

    # Card 1: Instalação e Armazenamento Local
    c1 = add_card(slide3, Inches(0.8), Inches(1.8), Inches(3.7), Inches(4.9), LOCAR_CARD, LOCAR_YELLOW)
    tb1 = slide3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.3), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "📱 Aplicativo PWA Instalável"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf1.add_paragraph()
    p.text = (
        "• Instalação sem App Store: O inspetor acessa o link e toca em 'Adicionar à Tela de Início'.\n\n"
        "• Service Worker & Cache: O código, telas e fotos são gravados no celular para abrir instantaneamente mesmo em modo avião.\n\n"
        "• Banco Local Dual-Layer:\n"
        "  - IndexedDB: Fotos e evidências pesadas.\n"
        "  - LocalStorage: Dados de frota e filas."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # Card 2: Sem Internet durante a Inspeção
    c2 = add_card(slide3, Inches(4.8), Inches(1.8), Inches(3.7), Inches(4.9), LOCAR_SURFACE, LOCAR_BORDER)
    tb2 = slide3.shapes.add_textbox(Inches(5.0), Inches(2.0), Inches(3.3), Inches(4.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🔒 Zero Risco de Perda"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf2.add_paragraph()
    p.text = (
        "• Trabalho Ininterrupto no Pátio: O inspetor preenche o checklist, tira fotos e assina sem travar.\n\n"
        "• Fila de Contingência (Sync Queue): Se não houver rede, a vistoria entra em uma fila segura no celular.\n\n"
        "• Segurança Máxima: Bateria descarregada, desligamento ou fechamento do app não apagam os dados preenchidos."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # Card 3: Indicador Visual no Cabeçalho
    c3 = add_card(slide3, Inches(8.8), Inches(1.8), Inches(3.7), Inches(4.9), LOCAR_CARD, LOCAR_YELLOW)
    tb3 = slide3.shapes.add_textbox(Inches(9.0), Inches(2.0), Inches(3.3), Inches(4.5))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "🚦 Indicadores no Topo"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf3.add_paragraph()
    p.text = (
        "O inspetor visualiza a conectividade em tempo real no cabeçalho:\n\n"
        "• Appwrite OK (Verde):\n"
        "  Todos os laudos já estão salvos e sincronizados na nuvem.\n\n"
        "• X Pendente(s) (Amarelo):\n"
        "  Há laudos na fila aguardando término de transmissão.\n\n"
        "• Offline Pátio (Vermelho):\n"
        "  Celular sem sinal; dados protegidos na memória local."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 4: COMO E QUANDO SINCRONIZA COM A NUVEM
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "3. Sincronização com a Nuvem (Appwrite Cloud)", "AUTOMAÇÃO COMPLETA: O USUÁRIO NÃO PRECISA FAZER NADA")

    # Card Grande Esquerda
    c_auto = add_card(slide4, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.9), LOCAR_CARD, LOCAR_BORDER)
    tb_auto = slide4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(6.2), Inches(4.5))
    tf_auto = tb_auto.text_frame
    tf_auto.word_wrap = True

    p = tf_auto.paragraphs[0]
    p.text = "⚡ Quando a Sincronização Acontece?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf_auto.add_paragraph()
    p.text = (
        "1. Automático ao Recuperar Sinal (Wi-Fi ou 4G/5G):\n"
        "Assim que o inspetor se aproxima da oficina ou entra na cobertura de rede, o navegador detecta o evento 'online' instantaneamente e descarrega toda a fila de vistorias pendentes (flushSyncQueue) para a nuvem.\n\n"
        "2. Automático ao Reabrir o Aplicativo com Internet:\n"
        "Se o app foi fechado offline, nos primeiros 3 segundos após reabrir com conexão, uma rotina em segundo plano verifica pendências e envia tudo automaticamente.\n\n"
        "3. Em Tempo Real (se já houver sinal no pátio):\n"
        "Se o aparelho estiver conectado no momento de assinar o laudo, a gravação na nuvem ocorre no mesmo segundo."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(12)

    # Card Direita
    c_man = add_card(slide4, Inches(7.8), Inches(1.8), Inches(4.7), Inches(4.9), LOCAR_SURFACE, LOCAR_YELLOW)
    tb_man = slide4.shapes.add_textbox(Inches(8.1), Inches(2.0), Inches(4.1), Inches(4.5))
    tf_man = tb_man.text_frame
    tf_man.word_wrap = True

    p = tf_man.paragraphs[0]
    p.text = "❓ É necessário fazer algo?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf_man.add_paragraph()
    p.text = "NÃO! A sincronização é 100% transparente e automática."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW
    p.space_before = Pt(6)

    p = tf_man.add_paragraph()
    p.text = (
        "O inspetor não precisa se preocupar em apertar botões para sincronizar.\n\n"
        "👆 Sincronização Manual (Opcional):\n"
        "Se o inspetor desejar forçar o envio imediato, basta tocar no botão 'Appwrite OK' ou 'Pendente(s)' no cabeçalho.\n\n"
        "O sistema tentará a conexão na hora e exibirá um aviso em tela: 'X registros sincronizados com sucesso!'."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 5: REGRAS DE TRANSIÇÃO AUTOMÁTICA DE STATUS
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "4. Regra Operacional: Transições Automáticas de Status", "AUTOMAÇÃO DE STATUS DA FROTA NA FINALIZAÇÃO DA VISTORIA")

    # Card Esquerda: Reprovado -> Manutenção
    c_rep = add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), LOCAR_CARD, LOCAR_YELLOW)
    tb_rep = slide5.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf_rep = tb_rep.text_frame
    tf_rep.word_wrap = True
    p = tf_rep.paragraphs[0]
    p.text = "⛔ SE REPROVADO NA INSPEÇÃO"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf_rep.add_paragraph()
    p.text = "Status da Frota passa automaticamente para:"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    p = tf_rep.add_paragraph()
    p.text = "⚙️ MANUTENÇÃO (PCM)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW
    p.space_before = Pt(4)

    p = tf_rep.add_paragraph()
    p.text = (
        "• Bloqueio Compulsório: Equipamento fica impedido de sair do pátio.\n"
        "• S.S. Imediata: Gera a Solicitação de Serviço ao PCM com fotos e normas.\n"
        "• Suspensão de Reserva: Se estava reservado para cliente, a reserva é suspensa e anotada como retida para manutenção.\n"
        "• Notificação PCM: Disparo de e-mail padrão e WhatsApp para a oficina."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(12)

    # Card Direita: Aprovado -> Disponível
    c_apr = add_card(slide5, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.9), LOCAR_SURFACE, LOCAR_BORDER)
    tb_apr = slide5.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.5))
    tf_apr = tb_apr.text_frame
    tf_apr.word_wrap = True
    p = tf_apr.paragraphs[0]
    p.text = "✓ SE 100% APROVADO NA INSPEÇÃO"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf_apr.add_paragraph()
    p.text = "Status da Frota passa automaticamente para:"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    p = tf_apr.add_paragraph()
    p.text = "🟢 DISPONÍVEL NO PÁTIO"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(4)

    p = tf_apr.add_paragraph()
    p.text = (
        "• Liberação Comercial: Equipamento apto para nova locação no pátio.\n"
        "• Limpeza de Pendências: Reserva antiga e dados transitórios são liberados.\n"
        "• Baixa Técnica no PCM: Qualquer S.S. que estava em aberto é concluída no sistema com o laudo de aprovação.\n"
        "• Laudo Liberado: Emissão do laudo com selo de liberação, QR Code e Hash SHA-256."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 6: GESTÃO DE FROTA & DASHBOARD DE BETIM
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "5. Gestão de Frota & Métricas CMMS Engeman®", "PAINEL GERENCIAL DA FILIAL BETIM / MG")

    # Imagem à esquerda
    if os.path.exists(img_dash):
        slide6.shapes.add_picture(img_dash, Inches(0.8), Inches(1.8), width=Inches(6.4))

    # Card à direita
    c_dash = add_card(slide6, Inches(7.4), Inches(1.8), Inches(5.1), Inches(4.9), LOCAR_CARD, LOCAR_BORDER)
    tb_dash = slide6.shapes.add_textbox(Inches(7.65), Inches(2.0), Inches(4.6), Inches(4.5))
    tf_dash = tb_dash.text_frame
    tf_dash.word_wrap = True

    p = tf_dash.paragraphs[0]
    p.text = "Indicadores em Tempo Real"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf_dash.add_paragraph()
    p.text = (
        "• 272 Ativos Cadastrados (Bases Fleet 1 & 2):\n"
        "  - Divisão PTA: 245 plataformas elevatórias (Genie, JLG, Haulotte, Skyjack).\n"
        "  - Divisão Guindastes & Outros: 27 ativos (Liebherr, Grove, Muncks, Empilhadeiras).\n\n"
        "• Status Operacionais Oficiais da Filial Betim:\n"
        "  - 138 em Manutenção (com prazos exatos de saída e dias p/ liberar do PCM).\n"
        "  - 98 Locadas em operação em clientes vigentes.\n"
        "  - 33 Disponíveis no pátio aptas para mobilização imediata.\n"
        "  - 3 Reservadas comerciais em negociação.\n\n"
        "• Controle de Prazos de Saída da Manutenção:\n"
        "  Data de liberação e contagem regressiva visíveis em cada card e exportação CSV.\n\n"
        "• Alerta de Manutenção Preventiva:\n"
        "  Identifica automaticamente horímetros próximos do ciclo de 250h/500h com badge de preventiva."
    )
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 7: O PASSO A PASSO DA INSPEÇÃO TÉCNICA
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "6. O Processo de Inspeção Técnica no Pátio", "STEPPER GUIADO COM EVIDÊNCIAS E ASSINATURA TOUCH")

    if os.path.exists(img_insp):
        slide7.shapes.add_picture(img_insp, Inches(0.8), Inches(1.8), width=Inches(6.4))

    c_insp = add_card(slide7, Inches(7.4), Inches(1.8), Inches(5.1), Inches(4.9), LOCAR_CARD, LOCAR_YELLOW)
    tb_insp = slide7.shapes.add_textbox(Inches(7.65), Inches(2.0), Inches(4.6), Inches(4.5))
    tf_insp = tb_insp.text_frame
    tf_insp.word_wrap = True

    p = tf_insp.paragraphs[0]
    p.text = "Rigor do Checklist Normativo"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf_insp.add_paragraph()
    p.text = (
        "1. Identificação Completa:\n"
        "Matrícula do inspetor, telefone, turno e horímetro atual.\n\n"
        "2. Vistoria Setorizada:\n"
        "Itens divididos por Estrutura Externa, Operação, Dispositivos de Emergência e Mecânica/Hidráulica.\n\n"
        "3. Foto Forense Obrigatória:\n"
        "Qualquer item marcado como Não Conforme exige captura de foto com carimbo de TAG, data e hora.\n\n"
        "4. Assinatura Digital Touch no Canvas:\n"
        "Obrigatória antes de concluir. Sem assinatura, o laudo não é emitido."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 8: CENTRAL DE SERVIÇOS DO PCM E E-MAILS
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "7. Central do PCM & Ordens Corretivas", "PLANEJAMENTO, CONTROLE DE MANUTENÇÃO E NOTIFICAÇÕES")

    if os.path.exists(img_pcm):
        slide8.shapes.add_picture(img_pcm, Inches(0.8), Inches(1.8), width=Inches(6.4))

    c_pcm = add_card(slide8, Inches(7.4), Inches(1.8), Inches(5.1), Inches(4.9), LOCAR_CARD, LOCAR_BORDER)
    tb_pcm = slide8.shapes.add_textbox(Inches(7.65), Inches(2.0), Inches(4.6), Inches(4.5))
    tf_pcm = tb_pcm.text_frame
    tf_pcm.word_wrap = True

    p = tf_pcm.paragraphs[0]
    p.text = "Fluxo Integrado da Oficina"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = LOCAR_YELLOW

    p = tf_pcm.add_paragraph()
    p.text = (
        "• Solicitação de Serviço (S.S.) Instantânea:\n"
        "A cada máquina reprovada, gera-se uma S.S. oficial com a lista de defeitos e fotos anexadas.\n\n"
        "• Notificação Multicanal:\n"
        "  - E-mail corporativo em HTML com formato pericial.\n"
        "  - Botão de envio rápido para o WhatsApp do gestor do PCM.\n\n"
        "• Workflow de 4 Fases:\n"
        "  Aberta ➔ Em Planejamento ➔ Em Execução ➔ Concluída.\n\n"
        "• Reinspeção Mandatória:\n"
        "  Após o PCM concluir, o equipamento passa por nova inspeção para retornar a DISPONÍVEL."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 9: LAUDO PERICIAL & INTEGRIDADE FORENSE (SHA-256 / QR CODE)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "8. Laudo Técnico Pericial & Integridade Forense", "BLINDAGEM JURÍDICA E CERTIFICAÇÃO DIGITAL")

    if os.path.exists(img_laudos):
        slide9.shapes.add_picture(img_laudos, Inches(0.8), Inches(1.8), width=Inches(6.4))

    c_laudo = add_card(slide9, Inches(7.4), Inches(1.8), Inches(5.1), Inches(4.9), LOCAR_SURFACE, LOCAR_YELLOW)
    tb_laudo = slide9.shapes.add_textbox(Inches(7.65), Inches(2.0), Inches(4.6), Inches(4.5))
    tf_laudo = tb_laudo.text_frame
    tf_laudo.word_wrap = True

    p = tf_laudo.paragraphs[0]
    p.text = "Validade Pericial e Jurídica"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf_laudo.add_paragraph()
    p.text = (
        "• Hash Criptográfico SHA-256:\n"
        "Cada laudo emite uma assinatura matemática inviolável com base nos dados e fotos da vistoria.\n\n"
        "• QR Code de Validação em Campo:\n"
        "Auditores ou clientes em obra podem escanear o QR Code impresso e consultar a integridade original.\n\n"
        "• Exportação Direta para PDF:\n"
        "Emissão em 1 clique sem depender de impressoras ou conexões externas.\n\n"
        "• Histórico Perpétuo do Ativo:\n"
        "Rastreabilidade completa de todas as intervenções para auditorias do Ministério do Trabalho e clientes."
    )
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_LIGHT
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 10: CLASSIFICAÇÃO DE CLIENTES E MATRIZ DE RIGOR (PALETA LOCAR)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "9. Matriz de Rigor por Cliente (Client Tiers)", "DIRETRIZES TÉCNICAS E COMERCIAIS POR CLIENTE")

    tiers = [
        {"tier": "TIER AA", "title": "Grandes Minerações & Siderurgias", "clients": "Vale, Usiminas, Anglo American, CSN, ArcelorMittal", "req": "Rigor Máximo: Pintura 100% impecável, testes de carga, ART mecânica, horímetro estrito e vistoria pericial.", "highlight": True},
        {"tier": "TIER A", "title": "Contratadas dos Grandes Players", "clients": "Andrade Gutierrez, Camargo Corrêa, Manserv, Tenenge", "req": "Rigor Elevado: Atendimento integral às NRs 11/12/18 e homologação de requisitos da contratante principal.", "highlight": False},
        {"tier": "TIER B", "title": "Obras Públicas & Infraestrutura", "clients": "Construtoras de Rodovias, Galpões e Prefeituras", "req": "Rigor Moderado: Equipamento 100% funcional com itens mandatórios de segurança testados e laudo válido.", "highlight": False},
        {"tier": "TIER C", "title": "Locação Operacional Padrão", "clients": "Locações Avulsas, Indústrias Locais e Prestadores", "req": "Rigor Padrão: Conformidade mecânica, hidráulica e operacional plena antes da mobilização no pátio.", "highlight": False}
    ]

    for i, tr in enumerate(tiers):
        tx = Inches(0.8 + i * 2.98)
        ty = Inches(1.8)
        tw = Inches(2.8)
        th = Inches(4.8)
        
        # Uso estrito da paleta Locar: Cinza, Amarelo e Branco
        card_bg = LOCAR_SURFACE if tr["highlight"] else LOCAR_CARD
        card_border = LOCAR_YELLOW if tr["highlight"] else LOCAR_BORDER
        add_card(slide10, tx, ty, tw, th, card_bg, card_border)

        # Barra amarela Locar
        bar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, ty, tw, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = LOCAR_YELLOW
        bar.line.fill.background()

        tb = slide10.shapes.add_textbox(tx + Inches(0.2), ty + Inches(0.3), tw - Inches(0.4), th - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tr["tier"]
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = LOCAR_YELLOW

        p = tf.add_paragraph()
        p.text = tr["title"]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = f"Principais Clientes:\n{tr['clients']}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

        p = tf.add_paragraph()
        p.text = f"Exigências Técnicas:\n{tr['req']}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 11: ENTRADA EM OPERAÇÃO & CARGA DE DADOS OFICIAIS
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "10. Entrada em Operação & Carga de Dados Oficiais", "TRANSIÇÃO DA FASE DE TESTES PARA OPERAÇÃO REAL • PURGA DE DADOS E CARGA CORPORATIVA")

    rollout_pillars = [
        {
            "icon": "🧹",
            "title": "Purga Rigorosa de Testes",
            "desc": "Remoção integral de todos os laudos simulados, vistorias preliminares, solicitações fictícias do PCM e usuários de homologação. Elimina contaminação de dados."
        },
        {
            "icon": "👥",
            "title": "Governança & Cadastro por CPF / E-mail (RBAC)",
            "desc": "Perfil 'admin' exclusivo do gestor Maikon Pinho (maikon.pinho@locar.com.br). Cadastro oficial de operadores habilitado por CPF ou E-mail corporativo, com PIN com hash SHA-256 e controle soberano de permissões."
        },
        {
            "icon": "🚜",
            "title": "Atualização de Status da Frota",
            "desc": "Atualização em lote do status real de operação dos 272 ativos da base Betim: 138 em manutenção (com prazos exatos de saída do PCM), 98 locadas, 33 disponíveis e 3 reservadas."
        },
        {
            "icon": "📢",
            "title": "Informativo de Falhas & Sugestões (E-mail Direto)",
            "desc": "Canal no cabeçalho e rodapé para relato de bugs e sugestões de melhoria. Cada chamado gera protocolo único e é despachado por e-mail formal para maikon.pinho@locar.com.br com inbox de gestão no painel."
        }
    ]


    for i, p_item in enumerate(rollout_pillars):
        bx = Inches(0.8 + (i % 2) * 6.0)
        by = Inches(1.8 + (i // 2) * 2.5)
        add_card(slide11, bx, by, Inches(5.7), Inches(2.2), LOCAR_CARD, LOCAR_BORDER)

        tb = slide11.shapes.add_textbox(bx + Inches(0.3), by + Inches(0.25), Inches(5.1), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{p_item['icon']}  {p_item['title']}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = LOCAR_YELLOW

        p = tf.add_paragraph()
        p.text = p_item["desc"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 12: OTIMIZAÇÃO RESPONSIVA — VISÃO COMPUTADOR & MOBILE (PWA)
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "11. Otimização Responsiva: Computador & Mobile (PWA)", "AUTO-ADAPTAÇÃO ERGONÔMICA PARA DESKTOP NO ESCRITÓRIO E SMARTPHONES NO PÁTIO")

    responsive_features = [
        {
            "icon": "💻",
            "title": "Visão Computador & Escritório (Desktop)",
            "desc": "Visão analítica ampla com grids em múltiplas colunas (1440px), painel de KPIs executivos com status por divisão (PTA / Guindastes), visualização ampla de laudos e atalhos de gestão."
        },
        {
            "icon": "📱",
            "title": "Visão Celular & Tablet no Pátio (Mobile)",
            "desc": "Adaptação automática para smartphones (Android e iOS): layout em coluna única, botões e touch-targets de fácil acionamento com o polegar (>=44px) e abas de rolagem tátil suave."
        },
        {
            "icon": "📷",
            "title": "Câmera e Assinatura Touch em Campo",
            "desc": "Acionamento da câmera traseira física do celular (facingMode: environment) para fotos de avarias, além de canvas de assinatura digital por toque com bloqueio de scroll (touch-action: none)."
        },
        {
            "icon": "⚡",
            "title": "PWA Instalável & IA Copilot Flutuante",
            "desc": "Pode ser adicionado à tela inicial do celular como App nativo (PWA com Service Worker offline). O balão de IA adapta seu tamanho dinamicamente em qualquer resolução de tela."
        }
    ]

    for i, rf in enumerate(responsive_features):
        bx = Inches(0.8 + (i % 2) * 6.0)
        by = Inches(1.8 + (i // 2) * 2.5)
        add_card(slide12, bx, by, Inches(5.7), Inches(2.2), LOCAR_CARD, LOCAR_BORDER)

        tb = slide12.shapes.add_textbox(bx + Inches(0.3), by + Inches(0.25), Inches(5.1), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{rf['icon']}  {rf['title']}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = LOCAR_YELLOW

        p = tf.add_paragraph()
        p.text = rf["desc"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 13: CONCLUSÃO E BENEFÍCIOS CORPORATIVOS
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "12. Conclusão & Benefícios Estratégicos", "GANHOS TÉCNICOS, OPERACIONAIS E JURÍDICOS PARA A LOCAR")

    benefits = [
        {"icon": "⚡", "title": "Zero Papel no Pátio", "desc": "Processo 100% digitalizado em smartphones ou tablets, eliminando pranchetas e retrabalho de digitação."},
        {"icon": "🛡️", "title": "Segurança e Blindagem Jurídica", "desc": "Comprovação forense de vistoria técnica prévia rigorosa para Ministério do Trabalho e seguradoras."},
        {"icon": "🤝", "title": "Sinergia Pátio, PCM e Comercial", "desc": "Visibilidade instantânea sobre quais equipamentos estão aptos, reservados ou em reparos na oficina."},
        {"icon": "🌐", "title": "Resiliência Máxima Offline-First", "desc": "Operação garantida no pátio de Betim sem interrupções e sincronização automática transparente com a nuvem."}
    ]

    for i, b in enumerate(benefits):
        bx = Inches(0.8 + (i % 2) * 6.0)
        by = Inches(1.8 + (i // 2) * 2.5)
        add_card(slide13, bx, by, Inches(5.7), Inches(2.2), LOCAR_CARD, LOCAR_BORDER)

        tb = slide13.shapes.add_textbox(bx + Inches(0.3), by + Inches(0.25), Inches(5.1), Inches(1.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{b['icon']}  {b['title']}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = LOCAR_YELLOW

        p = tf.add_paragraph()
        p.text = b["desc"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(8)


    # Salva na raiz do projeto com o novo nome oficial
    output_path = os.path.join(base_dir, 'Sistema de Inspeção & Qualidade de Frota.pptx')
    prs.save(output_path)
    print(f"Sucesso: Apresentacao criada com as cores oficiais Locar em: {output_path}")

if __name__ == '__main__':
    create_presentation()
